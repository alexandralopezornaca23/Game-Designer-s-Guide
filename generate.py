#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Builds the documentation site from GameDesignerGuide.xlsx.

Single source of truth: the spreadsheet. This script reads every sheet, detects
its columns from the header row (by name, not by position, so it keeps working
when columns are added or moved), and writes the docs/ folder that MkDocs turns
into the site.

The site chrome is in English; the explanations in each entry stay in Spanish,
which is how they were written.

Usage:   python generate.py [path/to/GameDesignerGuide.xlsx]
"""
import os, re, sys, shutil, unicodedata, warnings
from collections import OrderedDict
import openpyxl

warnings.simplefilter("ignore")

ROOT = os.path.dirname(os.path.abspath(__file__))
XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "GameDesignerGuide.xlsx")
DOCS = os.path.join(ROOT, "docs")

HEADER_ROW = 3
FIRST_DATA_ROW = 4

# Sheets that are navigation or spreadsheet plumbing, not content.
SKIP_SHEETS = {"📌 Índice", "➕ Plantilla", "Listas"}

# Sheet name -> (display name, blurb). Order here is the order in the site.
PROGRAMS = OrderedDict([
    ("Blender",              ("Blender",       "Modelling, sculpting, UVs, nodes and animation.")),
    ("ZBrush",               ("ZBrush",        "Digital sculpting: all 25 top-menu palettes, control by control.")),
    ("Maya",                 ("Maya",          "Modelling and animation.")),
    ("Krita",                ("Krita",         "Digital painting, textures and 2D art.")),
    ("Photoshop",            ("Photoshop",     "Retouching and textures.")),
    ("Illustrator",          ("Illustrator",   "Vectors and UI art.")),
    ("Unity",                ("Unity",         "Game engine.")),
    ("Unreal Engine",        ("Unreal Engine", "Game engine.")),
    ("PureRef",              ("PureRef",       "Reference image boards.")),
    ("💡 Consejos Generales", ("General Tips",  "Workflow, organisation and portfolio practices.")),
    ("Python",               ("Python",        "Scripts for Blender (bpy) and Maya (maya.cmds).")),
    ("MEL",                  ("MEL",           "Scripts in Maya's native language.")),
    ("ZScript",              ("ZScript",       "ZBrush's own macro and automation language.")),
])

# Sheets whose content is code: published as code blocks, not tables.
CODE_SHEETS = {"Python", "MEL", "ZScript"}

# --------------------------------------------------------------------------- category names
# Category names live in the spreadsheet and many are written in Spanish. These
# maps translate them for the site's navigation. Anything not listed falls back
# to the original text, so adding a new category never breaks the build.
CATEGORY_PREFIXES = [
    ("Barra T: ",                     "T Panel: "),
    ("Menú File: ",                   "File Menu: "),
    ("Menú superior > ",              "Top Menu > "),
    ("Selector de tipo de editor: ",  "Editor Type Selector: "),
]
CATEGORY_SUFFIXES = [
    (" > Cabecera", " > Header"),
    (" > Panel N",  " > N Panel"),
    (" > Tabla",    " > Table"),
]
CATEGORY_EXACT = {
    "Ajustes": "Settings",
    "Animación": "Animation",
    "Archivo": "File",
    "Archivo y edición": "File & Editing",
    "Atajos esenciales": "Essential Shortcuts",
    "Añadir Workspace (botón +)": "Add Workspace (+ button)",
    "Barra T: Comunes a todos los modos": "T Panel: Common to all modes",
    "Barra T: Edit Mode (malla)": "T Panel: Edit Mode (mesh)",
    "Caja de herramientas": "Toolbox",
    "Capas": "Layers",
    "Control de versiones": "Version Control",
    "Copias de seguridad": "Backups",
    "Dibujo y anotación": "Drawing & Annotation",
    "Edición": "Editing",
    "Filtros de Selección": "Selection Filters",
    "Flujo de trabajo": "Workflow",
    "Flujos de trabajo": "Workflows",
    "Geometría": "Geometry",
    "Grupos": "Groups",
    "Herramientas": "Tools",
    "Interfaz": "Interface",
    "Interfaz - Ajustes": "Interface — Settings",
    "Interfaz - Barras": "Interface — Bars",
    "Interfaz - Paletas": "Interface — Palettes",
    "Interfaz - Ventana": "Interface — Window",
    "Jerarquía / GameObjects": "Hierarchy / GameObjects",
    "Menú Edit": "Edit Menu",
    "Modelado": "Modelling",
    "Máscara": "Masking",
    "Navegación": "Navigation",
    "Navegación viewport": "Viewport Navigation",
    "Navegación y vista": "Navigation & View",
    "Notas": "Notes",
    "Objetos": "Objects",
    "Organización de archivos": "File Organisation",
    "Personalizar atajos": "Customising Shortcuts",
    "Pinceles": "Brushes",
    "Pinceles / Color": "Brushes / Color",
    "Portafolio": "Portfolio",
    "Portapapeles": "Clipboard",
    "Productividad general": "General Productivity",
    "Reproducción": "Playback",
    "Selección": "Selection",
    "Simetría": "Symmetry",
    "Spreadsheet > Panel de datos": "Spreadsheet > Data Panel",
    "Tiempo": "Time",
    "Tool > Botones principales": "Tool > Main Buttons",
    "Tool > Sub-paletas": "Tool > Sub-palettes",
    "Transformar": "Transform",
    "Utilidades": "Utilities",
    "Ventanas": "Windows",
    "Visibilidad": "Visibility",
    "Vista / Zoom": "View / Zoom",
    "Vista de escena": "Scene View",
    "Visualización": "Display",
    "Workspaces (pestañas superiores)": "Workspaces (top tabs)",
    "Zscript": "ZScript",
}
LEVELS = {"Básico": "Basic", "Intermedio": "Intermediate", "Avanzado": "Advanced"}


def translate_category(name):
    if name in CATEGORY_EXACT:
        return CATEGORY_EXACT[name]
    out = name
    for es, en in CATEGORY_PREFIXES:
        if out.startswith(es):
            out = en + out[len(es):]
            break
    for es, en in CATEGORY_SUFFIXES:
        if out.endswith(es):
            out = out[: -len(es)] + en
            break
    return out


MKDOCS_TEMPLATE = """# GENERATED BY generate.py — do not hand-edit the nav: section
# (the rest of the file can be edited, but it is overwritten on regeneration)
site_name: Game Art Tools Reference
site_description: >-
  Shortcuts, tools and interface reference for Blender, ZBrush, Maya, Krita,
  Unity, Unreal Engine and more — searchable across every program at once.
site_author: Alex
copyright: © Alex · Game Design & Development

theme:
  name: material
  language: en
  icon:
    logo: material/palette-swatch-outline
  favicon: assets/favicon.png
  font:
    text: Inter
    code: JetBrains Mono
  features:
    - navigation.tabs
    - navigation.tabs.sticky
    - navigation.top
    - navigation.indexes
    - navigation.instant
    - navigation.tracking
    - toc.follow
    - search.suggest
    - search.highlight
    - search.share
    - content.tooltips
  palette:
    - media: "(prefers-color-scheme: light)"
      scheme: default
      primary: deep purple
      accent: pink
      toggle:
        icon: material/weather-night
        name: Switch to dark mode
    - media: "(prefers-color-scheme: dark)"
      scheme: slate
      primary: deep purple
      accent: pink
      toggle:
        icon: material/weather-sunny
        name: Switch to light mode

markdown_extensions:
  - admonition
  - attr_list
  - md_in_html
  - tables
  - toc:
      permalink: true
  - pymdownx.highlight:
      anchor_linenums: true
  - pymdownx.superfences
  - pymdownx.keys
  - pymdownx.details

plugins:
  # The interface is English, the entries are written in Spanish, so the search
  # index is built for both.
  - search:
      lang:
        - en
        - es

extra_css:
  - assets/extra.css

nav:
{nav}
"""


# --------------------------------------------------------------------------- helpers
def slug(text):
    """'Tool > Sub-palettes' -> 'tool-sub-palettes'."""
    t = unicodedata.normalize("NFKD", str(text))
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = t.lower().replace("&", " and ")
    t = re.sub(r"[^a-z0-9]+", "-", t).strip("-")
    return t or "uncategorised"


def clean(v):
    return "" if v is None else str(v).strip()


def cell(v):
    """Prepare text for a Markdown table cell."""
    t = clean(v).replace("|", "\\|")
    return re.sub(r"\s*\n\s*", " ", t)


def keys(v):
    """Key combos become <kbd> elements; menu paths stay italic."""
    t = cell(v)
    if not t or t == "—":
        return "—"
    if ">" not in t and len(t) < 46 and not re.search(r"[a-záéíóúñ]{6,}", t.lower()):
        parts = [p.strip() for p in re.split(r"\s*\|\s*", t)]
        out = []
        for p in parts:
            out.append(" + ".join(f"<kbd>{k.strip()}</kbd>" for k in p.split("+") if k.strip()))
        return " · ".join(out)
    return f"*{t}*"


def header_map(ws):
    m = {}
    for c in range(1, ws.max_column + 1):
        v = clean(ws.cell(row=HEADER_ROW, column=c).value)
        if v and v not in m:
            m[v] = c
    return m


def find_col(m, *names):
    for n in names:
        if n in m:
            return m[n]
    return None


def read_sheet(ws):
    m = header_map(ws)
    cols = {
        "icon":    find_col(m, "Icono Aprox."),
        "cat":     find_col(m, "Categoría"),
        "palette": find_col(m, "Paleta"),
        "action":  find_col(m, "Acción", "Nombre del script"),
        "where":   find_col(m, "Atajo / Ubicación", "Referencia"),
        "program": find_col(m, "Programa"),
        "code":    find_col(m, "Código"),
        "level":   find_col(m, "Nivel"),
        "notes":   find_col(m, "Consejo / Nota", "Descripción"),
        "fav":     find_col(m, "Favorito"),
    }
    rows = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        d = {k: (clean(ws.cell(row=r, column=i).value) if i else "") for k, i in cols.items()}
        if d["action"] or d["notes"]:
            d["level"] = LEVELS.get(d["level"], d["level"])
            rows.append(d)
    return rows


def write(path, text):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)


# --------------------------------------------------------------------------- pages
def table_page(title, rows, has_palette):
    out = [f"# {title}\n"]
    groups = OrderedDict()
    for r in rows:
        key = r["palette"] if (has_palette and r["palette"] and r["palette"] != "—") else ""
        groups.setdefault(key, []).append(r)

    for key, group in groups.items():
        if key and key.strip() != title.strip():
            out.append(f"\n## {key}\n")
        out.append("\n| | Action | Shortcut / Location | Level | Notes |")
        out.append("|---|---|---|---|---|")
        for r in group:
            star = " ⭐" if r["fav"] and r["fav"] not in ("-", "—") else ""
            out.append("| {} | **{}**{} | {} | {} | {} |".format(
                r["icon"] or "", cell(r["action"]), star, keys(r["where"]),
                cell(r["level"]) or "—", cell(r["notes"]) or ""))
        out.append("")
    return "\n".join(out) + "\n"


def code_page(title, rows):
    out = [f"# {title}\n"]
    current = None
    for r in rows:
        cat = translate_category(r["cat"]) if r["cat"] else ""
        if cat and cat != current:
            current = cat
            out.append(f"\n## {current}\n")
        out.append(f"\n### {r['action']}\n")
        meta = [x for x in (r["program"], r["level"]) if x]
        if meta:
            out.append("*" + " · ".join(meta) + "*\n")
        if r["notes"]:
            out.append(f"{r['notes']}\n")
        if r["code"]:
            out.append("```python")
            out.append(r["code"])
            out.append("```\n")
    return "\n".join(out) + "\n"


# --------------------------------------------------------------------------- main
def main():
    if not os.path.exists(XLSX):
        sys.exit(f"Spreadsheet not found: {XLSX}")

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    if os.path.isdir(DOCS):
        shutil.rmtree(DOCS)
    os.makedirs(DOCS)

    nav, summary, total = [], [], 0

    for sheet, (name, blurb) in PROGRAMS.items():
        if sheet not in wb.sheetnames or sheet in SKIP_SHEETS:
            continue
        rows = read_sheet(wb[sheet])
        if not rows:
            continue
        total += len(rows)
        folder = slug(name)
        has_palette = any(r["palette"] for r in rows)

        categories = OrderedDict()
        for r in rows:
            categories.setdefault(translate_category(r["cat"] or "Uncategorised"), []).append(r)

        # --- program index
        idx = [f"# {name}\n", f"{blurb}\n",
               f"**{len(rows)} entries** across {len(categories)} "
               f"{'category' if len(categories) == 1 else 'categories'}.\n",
               "\n| Category | Entries |", "|---|---:|"]
        for c, g in categories.items():
            idx.append(f"| [{c}]({slug(c)}.md) | {len(g)} |")
        write(os.path.join(DOCS, folder, "index.md"), "\n".join(idx) + "\n")

        children = [f"{folder}/index.md"]
        nav_groups = OrderedDict()
        for c, g in categories.items():
            text = code_page(c, g) if sheet in CODE_SHEETS else table_page(c, g, has_palette)
            write(os.path.join(DOCS, folder, f"{slug(c)}.md"), text)
            # Categories like "Properties > Render" or "Modifier: Mirror" are
            # grouped by their prefix so the sidebar stays navigable.
            m = re.match(r"^(.{2,28}?)\s*(?:>|:)\s+(.+)$", c)
            if m:
                nav_groups.setdefault(m.group(1).strip(), []).append(
                    {m.group(2).strip(): f"{folder}/{slug(c)}.md"})
            else:
                nav_groups.setdefault(None, []).append({c: f"{folder}/{slug(c)}.md"})

        for g, items in nav_groups.items():
            if g is None:
                children.extend(items)
            elif len(items) == 1:
                children.append({f"{g} > {list(items[0])[0]}": list(items[0].values())[0]})
            else:
                children.append({g: items})

        nav.append({name: children})
        summary.append((name, len(rows), len(categories), folder, blurb))

    # --- home page
    home = [
        "# Game Art Tools Reference\n",
        "Shortcuts, tools and interface reference for the programs I use to make "
        "games. It started as a personal spreadsheet while studying Game Design & "
        f"Development, and has grown into a catalogue of "
        f"**{format(total, ',')} entries** across {len(summary)} programs and "
        "scripting languages.\n",
        '!!! note "Heads up — the entries are written in Spanish"\n',
        "    The interface, navigation and structure of this site are in English, "
        "but each entry's explanation is in Spanish, which is how I wrote them. "
        "Tool, panel and menu names are always in English, exactly as they appear "
        "on screen.\n",
        '!!! tip "How to use it"\n',
        "    Use the **search box** at the top (or press <kbd>S</kbd>): it searches "
        "every program at once, which is the one thing a spreadsheet cannot do. If "
        "you already know where to look, pick the program from the tabs.\n",
        "## The programs\n",
        "\n| Program | Entries | Categories | What it covers |", "|---|---:|---:|---|",
    ]
    for n, nr, nc, folder, blurb in summary:
        home.append(f"| [{n}]({folder}/index.md) | {nr} | {nc} | {blurb} |")
    home += [
        "\n## How it is built\n",
        "The content lives in a single spreadsheet, which is where I edit it. A "
        "script (`generate.py`) reads that file, detects the columns from their "
        "header row and writes the Markdown for the whole site; MkDocs turns it "
        "into these pages and a GitHub Action publishes them on every push.\n",
        "It is the same *docs-as-code* approach studios use to document tools and "
        "pipelines: one source of truth, everything else generated.\n",
        "```\nGameDesignerGuide.xlsx  ──►  generate.py  ──►  docs/*.md  ──►  mkdocs build  ──►  site/\n```\n",
        "## Conventions\n",
        "- **Level** — Basic, Intermediate or Advanced, so you can tell everyday "
        "tools from once-a-month ones.\n"
        "- **English tool names** — buttons, palettes and menus are written exactly "
        "as they appear on screen; the explanations are in Spanish.\n"
        "- **⭐** — entries marked as favourites.\n"
        "- **PENDIENTE** — Spanish for *pending*: something still to be confirmed "
        "against the program itself. Where it appears, it is deliberate — I would "
        "rather flag a gap than invent an answer.\n",
    ]
    write(os.path.join(DOCS, "index.md"), "\n".join(home) + "\n")

    # --- mkdocs.yml with the navigation resolved
    def yaml_nav(items, indent=2):
        sp = " " * indent
        out = []
        for it in items:
            if isinstance(it, str):
                out.append(f"{sp}- {it}")
            else:
                (k, v), = it.items()
                k = k.replace('"', "'")
                if isinstance(v, str):
                    out.append(f'{sp}- "{k}": {v}')
                else:
                    out.append(f'{sp}- "{k}":')
                    out.extend(yaml_nav(v, indent + 4))
        return out

    write(os.path.join(ROOT, "mkdocs.yml"),
          MKDOCS_TEMPLATE.format(nav="\n".join(yaml_nav([{"Home": "index.md"}] + nav))))

    print(f"Generated {total} entries from {len(summary)} programs.")
    for n, nr, nc, _, _ in summary:
        print(f"   {n:20s} {nr:5d} entries · {nc:3d} categories")


if __name__ == "__main__":
    main()
